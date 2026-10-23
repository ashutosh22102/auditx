"""Build STUDY_EVIDENCE_AUDIT.csv/.xlsx and MISSING_SOURCES.csv.

Usage: python3 -I build_study_audit.py <thesis_dir> <audit_output_dir>
Verification statuses follow the five categories requested by the user. No original
full text was available, so no row can be 'Verified'.
"""
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

root = Path(sys.argv[1])
out = Path(sys.argv[2])
data = json.loads((root / "review_data/included_final.json").read_text(encoding="utf-8"))
recs = {}
for line in (root / "review_data/records.jsonl").read_text(encoding="utf-8").splitlines():
    if line.strip():
        r = json.loads(line)
        recs.setdefault(r["id"], r)
qc = list(csv.DictReader(open(out / "QUOTE_CHECK.csv", encoding="utf-8")))
qc_by = defaultdict(list)
for q in qc:
    qc_by[q["study_id"]].append(q)

# Bibkeys used in the thesis for included studies that have no bibkey in the data.
EXTRA_KEYS = {"W7161124699": "karami2026", "W4404364980": "siaw2024", "W4404181825": "feroz2024",
              "W7155097010": "maune2026", "W4367173429": "kemp2024", "W4297818236": "hossain2022"}

# ---- map citations to thesis sections -------------------------------------------------
sec_re = re.compile(r"\\(chapter|section|subsection)\{([^}]*)\}")
cite_re = re.compile(r"\\(?:paren|text)?cite[s]?\*?(?:\[[^\]]*\])*\{([^}]*)\}")
key_secs = defaultdict(list)
chap_no = {"Chapter1": 1, "Chapter2": 2, "Chapter3": 3, "Chapter4": 4, "Chapter5": 5}
for f in sorted((root / "Chapters").glob("Chapter*.tex")):
    c = chap_no[f.stem]
    s = ss = 0
    label = "Ch. %d" % c
    for ln, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
        if line.lstrip().startswith("%"):
            continue
        m = sec_re.search(line)
        if m:
            if m.group(1) == "section":
                s += 1
                ss = 0
                label = "%d.%d %s" % (c, s, m.group(2))
            elif m.group(1) == "subsection":
                ss += 1
                label = "%d.%d.%d %s" % (c, s, ss, m.group(2))
        for cm in cite_re.finditer(line):
            for k in cm.group(1).split(","):
                k = k.strip()
                tag = "%s (line %d)" % (label, ln)
                if not key_secs[k] or not key_secs[k][-1].startswith(label):
                    key_secs[k].append(tag)

# ---- thesis claims for cited studies (paraphrased from the LaTeX source) -----------------
CLAIM = {
    "firmansyah2025": "Only H1-direct study from a 'digital business context'; path coefficient 0.361 AI capability->sustainable CA; n=70; low quality (4.2.1, 4.4.2).",
    "shao2026": "AI capability positively related to CA in 546 Chinese SMEs, 'partly mediated' by business model innovation (Ch.1, 2.3.2); H1 direct; H6 indirect; merges AI capability with DC (4.3.4).",
    "almereb2026": "Generative AI related to cost- and differentiation-based CA; H1 direct (Table 4.1).",
    "kim2026": "AI capabilities related to differentiation and low-cost CA in 400 Korean exporters; H1 direct.",
    "giang2026": "AI capability -> CA in 463 Vietnamese SMEs (jointly reported with other digital capabilities); H1 direct, H2/H6 indirect.",
    "gafsi2026": "AI capabilities -> 'competitive performance' in 120 Tunisian family firms; H1 direct.",
    "almamary2025": "AI capabilities -> CA and venture success, 327 entrepreneurs; H1 direct.",
    "tu2026": "AI capability -> competitiveness via learning and resilience, 304 Chinese manufacturers; H1 direct; H6 indirect.",
    "waqas2024": "AI capability (jointly with BDA and IoT capability) -> green CA, CPEC manufacturing; H1 direct.",
    "almheiri2024": "Same sample of 344 UAE public managers as almheiri2025; AI capability enhances DC; H2 direct (one independent sample).",
    "almheiri2025": "Same sample of 344 UAE public managers as almheiri2024; AI capability enhances DC; H2 direct.",
    "fano2025": "AI ambidexterity 'fostered sensing, seizing and transforming' in 102 European SMEs; H2 direct.",
    "neirafernandez2026": "AI readiness related to dynamic digital capabilities in Spanish firms; H2 direct.",
    "texeira2026": "High-quality; AI capability not significantly related to sustainable CA once AI-enabled sustainability capabilities included; conditional indirect 0.086->0.155 with strategic flexibility (4.2.1, 4.3.4, abstract-level claim 'in the strongest studies').",
    "wang2026": "High-quality time-lagged survey; AI capability no significant main effect on sustainable competitive performance (b=0.05); 'the only high-quality study that tests the relationship directly finds no significant effect' (4.2.1, 5.2).",
    "guo2025": "High-quality; complementors (cross-border platform sellers); AI adoption x strategic agility -> product innovation (0.519) and brand upgrading (0.336) advantages; IMC indirect 0.346/0.242.",
    "malik2021": "AI capability -> innovation capability (0.68) and organisational capability (0.62); indirect 0.31 and 0.27 on CA; moderate quality (4.2.2, 4.3.4, 5.3).",
    "srahin2026": "Augmented analytics adoption: total effect 0.542 on CA; direct effect small with CI including zero; indirect via agility 0.445 (4.2.1, 4.3.4).",
    "dzreke2025": "Only study measuring persistence; 1,200 listed firms; stronger effect in FinTech and e-commerce; reporting doubts noted (4.1, 4.4.1, 4.4.2).",
    "li2026": "AI orientation -> absorptive capacity (0.312) in '502 Chinese tourism firms'; indirect 0.067, 0.057, serial 0.009 on performance (4.2.2, 4.3.2-4.3.4).",
    "noor2026": "Only e-commerce study with explicit DC dimension: GenAI adoption -> market sensing (0.571); serial indirect 0.110; sole basis of H3 'partly supported'.",
    "deng2025": "DC modelled as antecedent of GenAI adoption (0.453); listed as one of four quantitative DC-dimension studies (4.3).",
    "adamu2026": "fsQCA, 268 Nigerian firms; AI capability and cloud infrastructure substitutes; counted as focal 'digitally active' context.",
    "mukhopadhyay2024": "Qualitative, 26 interviews; AI-enabled marketing practices 'strengthened dynamic capabilities' in Indian FMCG firms; moderate quality.",
    "quoc2026": "AI adoption intensity as moderator of DC->performance in Vietnamese start-ups.",
    "awamleh2022": "E-commerce enterprises, Jordan, 224 respondents; IT capabilities mediate AI->CA; H6 indirect.",
    "bunea2026": "Electronic retailing; AI implementation -> benefits and customer experience -> CA; H2/H6 indirect.",
    "mashi2024": "AI adoption not significantly related to CA but AI applications were; H1 mixed.",
    "chen2026": "AI infrastructure did not achieve CA directly, only via green innovation; listed among 'the strongest studies' where the direct effect disappears (4.4.1, 5.2).",
    "sofiyan2026": "No direct effect of AI capability on change capacity; only under transformational leadership; H2 contradictory.",
    "kokash2026": "Tested and did not find DC mediation of AI capability->performance; H6 contradictory.",
    "ayoub2026": "Measures 'AI-enabled dynamic capability', merging the constructs (4.3.4, 5.3).",
    "aldhaheri2024": "487 SME managers/owners; AI applications and DC mediate turbulence->performance (2.3.2, 4.4.1).",
    "alnofeli2025": "AI-powered CRM capabilities mapped onto DC micro-foundations (4.2.2).",
    "ghobakhloo2025": "Delphi of 14 experts mapped AI onto sensing/seizing/transforming; described as conceptual in 4.3.",
    "dzreke2025b": "Only study of platform owners; conceptual; uncertain peer-review status (4.4.2).",
    "li2025": "Ch.2: GenAI capabilities -> CA via service innovation under dynamism; Ch.4: dynamism 'shaped the role of AI'.",
    "hossain2022": "Ch.2: relates marketing analytics capability and AI adoption to CA in manufacturing (title-level claim).",
    "karami2026": "Conceptual: AI supports sensing, seizing and transforming (4.3).",
    "siaw2024": "Conceptual: AI may substitute for HI internally but be ineffective for external scanning (4.3.1).",
    "feroz2024": "Mixed methods: AI-related capabilities per DC dimension from expert interviews (4.3).",
    "maune2026": "Mixed methods, 312 respondents Zimbabwe; AI enhanced sensing via monitoring, real-time data, predictive analytics (4.3.1).",
    "kemp2024": "Conceptual: AI must be 'situated' to yield CA (Ch.1, 2, 4.4.1, 5.3).",
}

# ---- per-study audit findings ------------------------------------------------------------
# status, discrepancy, action
V2 = {
    "W7203683920": "Data link removed in review_data_corrected; thesis now reports 119 independent studies / 94 empirical samples; H2 indirect 41 -> 42 (resolved)",
    "gafsi2026": "Kept as H1 directly reported but marked provisional (D*) in Table 4.1 and 4.2.1; unresolved pending full text",
    "giang2026": "Marked provisional (D*) in Table 4.1; caveat added in 4.2.1",
    "waqas2024": "Marked provisional (D*) in Table 4.1; caveat added in 4.2.1",
    "almereb2026": "Marked provisional (D*) in Table 4.1; caveat added in 4.2.1",
    "almamary2025": "Marked provisional (D*) in Table 4.1; caveat added in 4.2.1",
    "firmansyah2025": "Marked D* and dagger (publication unconfirmed); context description qualified in 4.2.1 and 4.4.2; unresolved",
    "almheiri2024": "Table 4.1 now shows H1 I; Ch.2 wording changed to association",
    "almheiri2025": "Appendix C now cites via bibliography (2025), so text and appendix agree; same-sample status unchanged (pair)",
    "fano2025": "Wording changed to 'fostered dynamic capabilities, measured as sensing, seizing and transforming'; bib year 2026 vs OpenAlex 2025 unresolved (Appendix C now uses citation year)",
    "wang2026": "Described as a moderation model whose null main effect is not a dedicated H1 test (4.2.1, 5.2); classification kept as contradictory; unresolved pending full text",
    "texeira2026": "'Strategic flexibility' replaced by 'sustainable food innovation orientation' throughout (4.2.1, 4.3.4, 4.4.1, Fig. 4.4, 5.3, 5.4) (resolved at abstract level); null direct effect and 'high' rating unverified",
    "malik2021": "Dagger marker (publication details unconfirmed) in Table 4.1 and Appendix C; caveat in 4.2.2; unresolved",
    "chen2026": "Removed from 'strongest studies'; described as abstract-level study of AI infrastructure; added non-significant routes (resolved)",
    "srahin2026": "p = 0.036 reported alongside the CI; 'no remaining direct effect' replaced (resolved, full-text values unverified)",
    "guo2025": "Thesis now states direct and indirect paths both significant; complementor label added in 5.1/5.3.2 (resolved; estimates unverified)",
    "adamu2026": "Counted as 'closely related digital context', not e-commerce; abstract and summaries now say four empirical e-commerce/platform studies (resolved)",
    "deng2025": "Removed from list of quantitative DC-dimension studies; described as overall-DC antecedent study (resolved)",
    "ghobakhloo2025": "Described as Delphi study, not conceptual (resolved)",
    "mukhopadhyay2024": "Quality changed moderate -> n.a. (abstract only, Section 3.5 rule); moved within Table 4.1; all dependent counts recalculated; direction of claim reworded (resolved)",
    "li2026": "Indirect effects reassigned from H6 to H4 and H5 (assignment rule, consistent with H3); H4/H5 now partly supported; 'firms' -> 'managers'; marked E in Appendix C",
    "noor2026": "Dagger marker; sole basis of H3 (unchanged); unresolved pending full text",
    "li2025": "Removed from the boundary-condition claim in 4.4.1; marked T in Appendix C; eligibility unresolved",
    "siaw2024": "4.3.1 now states the abstract's inconsistency and that the proposition needs full-text checking (resolved wording; content unresolved)",
    "shao2026": "Ch.2 'had a positive effect ... partly mediated' -> association wording, 'mediator' (resolved)",
    "sofiyan2026": "Marked E in Appendix C; E5 interpretation documented in 3.2.2",
    "dzreke2025": "Dagger marker in Table 4.1 and Appendix C",
    "dzreke2025b": "Dagger marker in Appendix C",
    "Shore2024": "Marked E in Appendix C; E5 interpretation documented in 3.2.2",
    "Tian2023": "Marked E in Appendix C; E5 interpretation documented in 3.2.2",
    "W4407627725": "Appendix C author restored from OpenAlex record ('- et al.' -> 'Kabir et al.')",
}

NOTE = {
    "W7203683920": ("Mismatch",
        "Data field same_sample_as links this hotel study (400 hotel managers, UAE; per abstract) to the Almheiri public-manager sample (344). Abstracts show different samples. The stored link is what makes the '118 independent samples' and H2 'indirectly relevant = 41' figures reproduce.",
        "Remove same_sample_as link; recount independent samples (119 publications-level, 94 empirical samples) and H2 indirect count (42 under the same priority rule)."),
    "gafsi2026": ("Needs checking",
        "Revised in v2 (first audit said Mismatch): the abstract names the dependent variable 'competitive performance' (aim) but reports the finding as 'AI capabilities significantly enhance firm performance'; the extraction flags the CA classification as tentative. Whether the outcome is CA-type (like Wang 2026's competitive performance) or firm performance (indirectly relevant under the Outcome rule) cannot be settled from the abstract.",
        "Obtain full text and check the outcome measure; if it is firm performance, reclassify as H1 indirectly relevant (H1 direct 9 -> 8)."),
    "almheiri2025": ("Partially verified",
        "Same-sample claim supported at abstract level (both: 344 UAE public managers, 26-item AI capability scale). Year inconsistency: bibliography 2025 vs OpenAlex/Appendix C 2024, so Appendix C lists two 'Almheiri et al. (2024)' entries while text cites 2024 and 2025. Outcome is performance, public sector.",
        "Harmonise year in Appendix C with bibliography; obtain full text to confirm sample overlap and AI->DC estimate."),
    "almheiri2024": ("Partially verified",
        "Same-sample claim supported at abstract level. Thesis Table 4.1 omits the H1 'indirectly relevant' coding held in the data. Mediation is 'partial' per abstract; method unclear.",
        "Obtain full text; confirm sample overlap and estimates."),
    "fano2025": ("Partially verified",
        "Bibliography year 2026 vs cite key/data/Appendix C 2025 (text will print 2026). Abstract gives no dimension-specific results, so 'fostered sensing, seizing and transforming' (4.2.2) slightly overstates; outcome is SME performance.",
        "Fix year; phrase as 'fostered dynamic capabilities (measured as sensing, seizing, transforming)'; check full text for dimension-level paths."),
    "wang2026": ("Needs checking",
        "Abstract presents AI capabilities only as a moderator of SEO->disruptive innovation; the b=0.05 'main effect' is a coefficient in an interaction model read from a table the extraction calls 'garbled'; AIC and DI measured in the same wave. Treating this as the decisive high-quality test 'contradicting' H1 overstates what the study tested.",
        "Obtain full text; re-examine whether AI->SCP was hypothesised/tested; consider reclassifying from 'contradictory' to 'indirectly relevant' or describing as a conditional effect."),
    "texeira2026": ("Mismatch",
        "Thesis (4.3.4, 4.4.1, 4.4.3 Fig. 4.4, 5.3) states the indirect effect 'increased with strategic flexibility'; the source abstract names the moderator 'sustainable food innovation orientation' (the extraction's SFIO). Abstract otherwise confirms 400 restaurant decision-makers and bootstrapped indirect effects. The null direct effect (B=0.071, p=0.070) is full-text only and unchecked. Extraction notes a sample inconsistency (449 vs 803 administered) yet the study is rated 'high'.",
        "Replace 'strategic flexibility' with 'sustainable food innovation orientation' everywhere (incl. revised model boundary conditions); obtain full text to verify the null direct effect; reconsider 'high' rating."),
    "malik2021": ("Needs checking",
        "OpenAlex year 2021 conflicts with a Zenodo DOI (10.5281/zenodo.22806492) in a recent number range; extraction notes CB-SEM fit indices reported for a PLS-SEM model and reporting inconsistencies. Abstract supports beta 0.62/0.68 and indirect 0.31; 0.27 is full-text only. Used as key moderate-quality evidence for H2/H6 and managerial implications.",
        "Confirm publication year, volume/issue and peer-review status from the publisher site; verify estimates in full text; correct citation year."),
    "chen2026": ("Mismatch",
        "Thesis lists this among 'the strongest studies' in which the direct effect disappears (4.4.1, 5.2, abstract), but its quality is 'not assessable' (abstract only), the AI construct is AI infrastructure, and the abstract states that two of the routes from AI infrastructure were non-significant.",
        "Remove from 'strongest studies' wording; describe as abstract-level evidence on AI infrastructure."),
    "srahin2026": ("Partially verified",
        "Abstract supports 'relationship ... operates mainly through firm agility'. Extraction shows the direct path had t=2.108, p=0.036 while the BCa CI included zero; thesis reports only the CI. Respondents are individuals across sectors (not firms).",
        "Obtain full text; report the conflicting p-value and CI; avoid 'disappears'."),
    "guo2025": ("Partially verified",
        "Abstract (and a web search of the publisher/repository record) confirms 322 Chinese SMEs on cross-border platforms, coupling -> product innovation and brand upgrading, IMC mediation. Estimates are full-text only. Extraction notes text 0.242 vs table .255; AI adoption items rate the platform's AI use and AI is bundled with strategic agility (itself a DC), so the predictor is not a firm AI capability.",
        "Obtain accepted manuscript (Dundee repository) and verify estimates; keep 'complementor' label wherever cited (incl. 5.3 'For platform firms')."),
    "firmansyah2025": ("Needs checking",
        "0.361 consistent with abstract ('36.1%'). Abstract contains a stray unrelated sentence fragment; no DOI; journal peer-review status unverified. AI capability is modelled as a mediator of KMC->SCA. Abstract does not mention e-commerce or digital business; the 'digital MSMEs' context comes from unverified full-text extraction, yet the study carries the H1 context-rule argument.",
        "Obtain full text; confirm sample is digital/e-commerce MSMEs and peer review; otherwise drop from focal-context count."),
    "adamu2026": ("Partially verified",
        "Abstract: firms in 'digitally active sectors' in Nigeria (FinTech, banking etc. per extraction), not e-commerce or platforms. Counting it among the 'seven' e-commerce/platform studies overstates focal evidence.",
        "Label as digital-sector, not e-commerce/platform; adjust 'seven' wording."),
    "deng2025": ("Mismatch",
        "Thesis 4.3 lists deng2025 among the four quantitative studies measuring DC dimensions; the data code its DC construct as 'overall' (dimension list but no dimension-level paths) and the ten dimension studies in the data are Wei 2025, Arowolo 2026, Siaw 2024, Karami 2026 (conceptual), Ghobakhloo 2025 (Delphi), Feroz 2024, Maune 2026 (mixed), Fano 2025, Neira-Fernandez 2026, Noor 2026 (quantitative). Extraction also notes inconsistent reliability values and dubious references while rating it moderate.",
        "Correct the 4.3 breakdown (3 quantitative, 4 conceptual, 1 Delphi, 2 mixed); re-check moderate rating."),
    "ghobakhloo2025": ("Partially verified",
        "Abstract confirms Delphi with 14 experts mapping AI/IoT/BDA to sensing/seizing/transforming. Thesis 4.3 calls it conceptual; data code it as empirical 'other' (Delphi). AI bundled with other technologies.",
        "Describe as Delphi study, not conceptual."),
    "mukhopadhyay2024": ("Needs checking",
        "Rated 'moderate' although fulltext_used=False (the only appraised study rated from the abstract), contrary to the thesis statement that appraisal needed full text. Abstract frames DC 'through AI technologies' enhancing marketing; thesis says AI-enabled marketing 'strengthened dynamic capabilities' (direction overstated).",
        "Obtain full text before keeping a quality rating; rephrase direction."),
    "li2026": ("Partially verified",
        "Abstract confirms 502 managerial respondents (thesis says '502 Chinese tourism firms'), CB-SEM, bootstrapping; estimates are full-text only. Outcome is firm performance and no DC construct is measured (AC/EOR are framed as seizing/transforming), so inclusion under E5 is borderline.",
        "Say 'respondents' not 'firms'; document why the study passes E5."),
    "noor2026": ("Partially verified",
        "Abstract confirms e-commerce firms in Pakistan, cross-sectional PLS-SEM, serial pathways via market sensing. Estimates (0.571, 0.110) full-text only; sample size not reported ('400-500'); no DOI; journal peer-review unverified. Sole basis of H3 'partly supported'.",
        "Obtain full text; verify estimates, sample and peer review."),
    "neirafernandez2026": ("Partially verified",
        "Abstract supports DDC mediation of AI readiness -> digital maturity. 'Directly reported' H2 rests on a quote that infers AIR->DDC from a mediation sentence; AI readiness is a borderline AI-capability construct; n not reported.",
        "Obtain full text; confirm AIR->DDC path estimate; reconsider 'directly reported'."),
    "giang2026": ("Partially verified",
        "Abstract reports results jointly for five capabilities; no AI-specific estimate. Counted as 'directly reported' for H1 although the AI-specific relationship is not visible (thesis acknowledges).",
        "Obtain full text for AI-specific path or reclassify as indirectly relevant."),
    "almamary2025": ("Partially verified",
        "Abstract quote verbatim; but CA appears as one of several AI-capability effect areas contributing to venture success; extraction notes the role of CA is ambiguous.",
        "Check full text to confirm CA is modelled as an outcome of AI capability."),
    "almereb2026": ("Partially verified",
        "Abstract quote verbatim. Generative AI is a mediator of GSCM->CA; construct re-coded from 'use' to 'capability' by the checker, which is what makes it 'directly reported'.",
        "Check measurement items; if GenAI use, reclassify as indirectly relevant."),
    "li2025": ("Source unavailable",
        "Included on title only (no abstract, no full text). Thesis Ch.4 states that environmental dynamism 'shaped the role of AI' in this study, which cannot be known from the title.",
        "Obtain paper; confirm eligibility; remove result claims until read."),
    "hossain2022": ("Source unavailable", "Included on title only (no abstract, no full text).", "Obtain paper; confirm eligibility."),
    "siaw2024": ("Needs checking",
        "Source abstract is internally inconsistent: Findings say substituting HI with AI suits external scanning; Limitations say AI may be ineffective for external scanning. Thesis 4.3.1 relies on the second reading only. Bib year 2025 vs OpenAlex 2024.",
        "Read full text and state which proposition the paper makes."),
    "maune2026": ("Partially verified",
        "Abstract confirms 312 senior managers in Zimbabwe, Gioia analysis, AI enhancing sensing. Data code mixed methods; AI capability -> CI effectiveness; low quality.",
        "Obtain full text for quote."),
    "kemp2024": ("Partially verified", "Abstract supports the situated-AI argument. Year: bib 2024 vs Appendix C 2023 (OpenAlex).", "Harmonise year."),
    "dzreke2025": ("Partially verified",
        "Abstract confirms panel of 1,200 firms (2015-2023), 58% mediation by orchestration capability. E-commerce comparison is a full-text quote. Extraction records serious credibility concerns (data timing, misread Sargan test). Thesis notes doubts.",
        "Keep caveats; verify in full text before citing in Discussion."),
    "dzreke2025b": ("Partially verified", "Abstract matches; unsupported quantitative claims (e.g. 71%) and uncertain peer review; conceptual.", "Keep as context only."),
    "sofiyan2026": ("Partially verified",
        "Abstract confirms no direct AIC->OCC effect; TL moderation; 81 naval units. Military/public setting; outcome performance, mediator change capacity: borderline under E5.",
        "Document E5 decision."),
    "kokash2026": ("Partially verified", "Abstract confirms DC did not mediate AI capability -> performance; sample size and design not reported.", "Obtain full text for n and test."),
    "quoc2026": ("Partially verified", "Abstract confirms AI adoption intensity moderates DC->performance (0.282).", "None beyond full-text check."),
    "awamleh2022": ("Partially verified", "Abstract confirms 224 individuals from Jordanian e-commerce enterprises; IT capabilities mediate. Mediation test not described; data code context 'other_ecommerce_digital'.", "Obtain full text."),
    "bunea2026": ("Partially verified", "Abstract confirms electronic retailing, PLS-SEM + cIPMA + NVivo; sample size not in abstract.", "Obtain full text for n."),
    "mashi2024": ("Partially verified", "Abstract supports adoption n.s. vs applications significant; sample reported inconsistently (315 vs 346).", "Note inconsistency."),
    "ayoub2026": ("Partially verified", "Abstract confirms AIDC construct, n=254 US managers, significant direct effect.", "None."),
    "aldhaheri2024": ("Partially verified", "Abstract confirms 487 managers/owners; AI and DC as mediators of turbulence. Thesis Ch.4 groups it with studies where turbulence 'shaped the role of AI' - consistent.", "None."),
    "alnofeli2025": ("Partially verified", "Abstract confirms mapping to DC micro-foundations (64 articles + 24 interviews). A second Alnofeli 2025 paper (IJIM) is also included; possible overlapping data not assessed.", "Check overlap between the two Alnofeli 2025 papers."),
    "tu2026": ("Partially verified", "Abstract confirms direct effect on competitiveness and serial mediation; 'competitiveness' treated as CA.", "None beyond full text."),
    "waqas2024": ("Partially verified", "Abstract states TMCs (AIC, BDAC, IoTC) jointly supported green CA; no AI-specific estimate. Country only 'CPEC'.", "Check AI-specific path."),
    "shao2026": ("Partially verified",
        "Abstract confirms 546 Chinese SMEs, positive influence on CA, BMI as mediator. 'Partly mediated' (Ch.2) is not stated in the abstract; causal wording ('effect') in Ch.1/2.",
        "Verify partial vs full mediation in full text; use association language."),
    "kim2026": ("Partially verified", "Korean abstract confirms 400 exporters (online panel, April 2026), AI capability raised both CA types.", "None beyond full text."),
    "feroz2024": ("Partially verified", "Abstract: SLR of 72 articles + 28 expert interviews; perceptual, acquisition and transforming capabilities.", "None."),
    "karami2026": ("Partially verified", "Abstract matches conceptual framing.", "None."),
    "Shore2024": ("Needs checking", "Outcome entrepreneurial resilience; no CA and DC only as framing (resilience as DC): borderline under E5. n=87 (France).", "Document E5 decision."),
    "Tian2023": ("Needs checking", "AI bundled with IoT ('datafication'); outcome SC performance; supply chain resilience framed as DC. Borderline under E5 yet coded H1 indirectly relevant.", "Document E5 decision."),
    "Nikzat2025": ("Needs checking", "Full text used but rated 'not assessable' (so not all 74 n.a. ratings are due to abstract-only access); no DOI; extraction notes reference list with many unrelated citations.", "Check publication legitimacy."),
}


def fmt_ref(x, r):
    au = r.get("authors") or []
    if len(au) > 3:
        a = ", ".join(au[:3]) + " et al."
    else:
        a = ", ".join(au) if au else "[authors missing in record]"
    return "%s (%s). %s. %s.%s" % (a, x["yr"], x["title"], x["journal"], (" https://doi.org/" + x["doi"]) if x["doi"] else "")


rows = []
missing = []
for x in data:
    r = recs[x["id"]]
    key = x.get("bibkey") or EXTRA_KEYS.get(x["id"], "")
    has_abs = bool((r.get("abstract") or "").strip())
    access = ("full text used by extraction (not available to this audit)" if x["fulltext_used"]
              else ("abstract only" if has_abs else "title only (no abstract, no full text)"))
    quotes = qc_by[x["id"]]
    n_abs = sum(q["quote_check"].startswith(("verbatim", "near-verbatim")) for q in quotes)
    first_q = quotes[0]["quote"] if quotes else ""
    secs = "; ".join(key_secs.get(key, [])) if key else ""
    if not secs:
        secs = "Not cited individually; counted in aggregate figures/tables of Ch. 4 and listed in Appendix C"
    # default status
    if not has_abs and not x["fulltext_used"]:
        st, disc, act = ("Source unavailable", "Included on title only: no abstract in OpenAlex record and no full text; no results extracted.",
                         "Obtain full text; confirm eligibility or remove; recount totals.")
    elif x["fulltext_used"]:
        st = "Partially verified" if has_abs else "Source unavailable"
        disc = "Extraction claims full-text use; %d of %d quotations are not in the abstract and could not be checked." % (len(quotes) - n_abs, len(quotes))
        act = "Obtain full text and check quotations, estimates and page numbers."
    else:
        st = "Partially verified"
        disc = "All %d extracted quotations found in the OpenAlex abstract; the full paper was not inspected (no methods, estimates or page numbers checkable)." % len(quotes) if quotes else "No results extracted."
        act = "Obtain full text to confirm design, measures and results."
    for k in (key, x["id"], "%s%s" % (x["first"], x["yr"])):
        if k in NOTE:
            st, disc, act = NOTE[k]
            break
    ev = []
    if quotes:
        ev.append("%d/%d quotations verbatim or near-verbatim in OpenAlex abstract" % (n_abs, len(quotes)))
    if has_abs:
        ev.append("OpenAlex abstract inspected")
    ev.append("original paper NOT inspected (no PDFs provided; publisher/OA hosts unreachable from this environment)")
    rows.append({
        "Study ID": x["id"],
        "Thesis cite key": key,
        "Full bibliographic reference": fmt_ref(x, r),
        "DOI or stable URL": ("https://doi.org/" + x["doi"]) if x["doi"] else (r.get("oa_landing") or "https://openalex.org/" + x["id"]),
        "Original source file available?": "no",
        "Full text or abstract only": access,
        "Study context and country": "%s; context=%s; country=%s" % (x["industry"], x["context"], x["country"]),
        "Sample size and unit of analysis": "%s; respondents: %s" % (x["sample_size"] if x["sample_size"] is not None else "not reported", x["respondents"]),
        "Research design and analytical method": "%s; %s" % (x["design"], x["analysis_method"]),
        "AI construct and its operationalisation": "%s: %s" % (x["ai_construct_type"], x["ai_construct_name"]),
        "Dynamic-capability construct or related mechanism": "%s %s; related: %s" % (x["dc_construct"], ",".join(x["dc_dimensions"]), "; ".join(x["related_mediators"]) or "-"),
        "Competitive-advantage or performance outcome": "%s: %s (persistence measured: %s)" % (x["outcome_type"], x["outcome_name"], x["persistence_measured"]),
        "Tier / quality / evidence categories (data)": "Tier %s; %s; %s" % (x["tier"], x["overall_quality"], json.dumps(x["categories"])),
        "Main findings reported in the thesis": CLAIM.get(key, "Not discussed individually"),
        "Extracted quotation and page number": (first_q[:300] + (" [+%d more, see QUOTE_CHECK.csv]" % (len(quotes) - 1) if len(quotes) > 1 else "")) + (" | page: not recorded in extraction" if first_q else ""),
        "Thesis section using the evidence": secs,
        "Verification status": st,
        "Evidence supporting the verification": "; ".join(ev),
        "Discrepancy identified": disc,
        "Recommended action": act,
        "Thesis correction (v2)": V2.get(key) or V2.get(x["id"]) or V2.get("%s%s" % (x["first"], x["yr"])) or
            ("Marked T (title only, provisional) in Appendix C; disclosed in 3.5 and 4.1; eligibility unresolved" if (not has_abs and not x["fulltext_used"]) else "No thesis change required"),
    })
    prio = "1-critical" if key in ("firmansyah2025", "shao2026", "almereb2026", "kim2026", "giang2026", "gafsi2026", "almamary2025",
                                   "tu2026", "waqas2024", "almheiri2024", "almheiri2025", "fano2025", "neirafernandez2026", "texeira2026",
                                   "wang2026", "guo2025", "srahin2026", "chen2026", "kokash2026", "noor2026", "malik2021", "li2026") \
        else ("2-high" if key and key in CLAIM else ("2-high" if x["context"] != "none" or (not has_abs and not x["fulltext_used"]) else "3-normal"))
    why = []
    if key in ("firmansyah2025", "shao2026", "almereb2026", "kim2026", "giang2026", "gafsi2026", "almamary2025", "tu2026", "waqas2024"):
        why.append("H1 directly reported evidence")
    if key in ("almheiri2024", "almheiri2025", "fano2025", "neirafernandez2026"):
        why.append("H2 directly reported evidence")
    if key in ("texeira2026", "srahin2026", "chen2026"):
        why.append("abstract/conclusion claim that direct effect disappears")
    if key in ("wang2026", "kokash2026", "chen2026"):
        why.append("contradictory evidence")
    if key in ("noor2026",):
        why.append("sole basis of H3")
    if key in ("malik2021", "guo2025", "li2026", "texeira2026", "srahin2026"):
        why.append("H2/H6 key estimates")
    if x["context"] != "none":
        why.append("e-commerce/platform focal context")
    if not has_abs and not x["fulltext_used"]:
        why.append("included on title only")
    if key and key in CLAIM and not why:
        why.append("cited individually in thesis")
    missing.append({
        "Priority": prio,
        "Study ID": x["id"], "Thesis cite key": key, "First author": x["first"], "Year": x["yr"], "Title": x["title"],
        "Journal": x["journal"], "DOI": x["doi"],
        "Why needed": "; ".join(why) or "included study; extraction unverified",
        "What is currently available": access,
        "Candidate OA PDF (from OpenAlex record, not accessed)": r.get("oa_pdf") or "",
        "Candidate landing page (from OpenAlex record, not accessed)": r.get("oa_landing") or "",
        "Access status in this audit": "not retrieved: network policy blocked doi.org/OpenAlex/publisher hosts",
    })

order = {"1-critical": 0, "2-high": 1, "3-normal": 2}
missing.sort(key=lambda m: (order[m["Priority"]], m["First author"]))

with open(out / "STUDY_EVIDENCE_AUDIT.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)
with open(out / "MISSING_SOURCES.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(missing[0].keys()))
    w.writeheader()
    w.writerows(missing)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    for title, tab in (("Study audit", rows), ("Quote check", qc), ("Missing sources", missing)):
        ws = wb.active if title == "Study audit" else wb.create_sheet(title)
        ws.title = title
        cols = list(tab[0].keys())
        ws.append(cols)
        for rr in tab:
            ws.append([rr[c] for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4A5560")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 28
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(out / "STUDY_EVIDENCE_AUDIT.xlsx")
except ImportError:
    print("openpyxl not available; XLSX skipped")

from collections import Counter
print("rows", len(rows))
print("status", Counter(r["Verification status"] for r in rows))
print("access", Counter(r["Full text or abstract only"] for r in rows))
print("missing priority", Counter(m["Priority"] for m in missing))
print("cited keys without section mapping", [r["Thesis cite key"] for r in rows if r["Thesis cite key"] and r["Thesis section using the evidence"].startswith("Not cited")])
